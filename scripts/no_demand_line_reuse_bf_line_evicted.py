#!/usr/bin/env python3
"""Extract demand-origin line-reuse histograms into an Excel workbook."""

import argparse
import os
import re

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
except ImportError:
    raise SystemExit(
        "The 'openpyxl' library is required. Install it with: pip install openpyxl"
    )


DEFAULT_RESULTS_DIR = "../results_bingo/pref_l1_l2/"
DEFAULT_OUTPUT_FILE = "../Excel_Output/aiml_bingo/demand_line_reuse_count.xlsx"
CACHE_LEVELS = ("L1D", "L2C", "LLC")
SHEET_NAMES = {"L1D": "L1D", "L2C": "L2", "LLC": "LLC"}
DEFAULT_WAYS = {"L1D": 20, "L2C": 10, "LLC": 4}
HISTOGRAM_LABEL = "DEMAND LINE REUSE COUNT"


def natural_sort_key(value):
    return [
        int(text) if text.isdigit() else text.lower()
        for text in re.split(r"([0-9]+)", value)
    ]


def load_cache_ways(cache_header):
    ways = DEFAULT_WAYS.copy()
    if not os.path.exists(cache_header):
        return ways

    with open(cache_header, "r", errors="ignore") as file:
        content = file.read()

    for cache_name in CACHE_LEVELS:
        match = re.search(rf"#define\s+{cache_name}_WAY\s+(\d+)", content)
        if match:
            ways[cache_name] = int(match.group(1))
    return ways


def parse_reuse_counts(histogram_text):
    return {
        int(reuse_count): int(num_lines)
        for reuse_count, num_lines in re.findall(r"(\d+)\s*:\s*(\d+)", histogram_text)
    }


def parse_file(filepath):
    with open(filepath, "r", errors="ignore") as file:
        content = file.read()

    reuse_counts = {}
    for cache_name in CACHE_LEVELS:
        match = re.search(
            rf"{cache_name}\s*{HISTOGRAM_LABEL}\s*:\s*([^\n\r]*)",
            content,
        )
        if match:
            reuse_counts[cache_name] = parse_reuse_counts(match.group(1))
    return reuse_counts


def collect_records(results_dir):
    records = []
    for root, _, files in os.walk(results_dir):
        for filename in sorted(files, key=natural_sort_key):
            filepath = os.path.join(root, filename)
            if not os.path.isfile(filepath):
                continue

            reuse_counts = parse_file(filepath)
            if reuse_counts:
                records.append({"trace": filename, "reuse_counts": reuse_counts})
    return records


def build_reuse_row(trace, counts, max_reuse_column):
    count_values = [counts.get(reuse_count, 0) for reuse_count in range(max_reuse_column + 1)]
    overflow = sum(value for reuse_count, value in counts.items() if reuse_count > max_reuse_column)
    total_lines = sum(count_values) + overflow
    percentages = [100.0 * value / total_lines if total_lines else 0.0 for value in count_values]
    overflow_percentage = 100.0 * overflow / total_lines if total_lines else 0.0
    return [trace] + count_values + [overflow] + percentages + [overflow_percentage]


def write_workbook(records, output_file, cache_ways):
    workbook = Workbook()
    workbook.remove(workbook.active)

    for cache_name in CACHE_LEVELS:
        max_reuse_column = cache_ways[cache_name]
        worksheet = workbook.create_sheet(SHEET_NAMES[cache_name])
        headers = ["Trace"] + [str(value) for value in range(max_reuse_column + 1)]
        headers += [f">{max_reuse_column}"]
        headers += [f"{value} %" for value in range(max_reuse_column + 1)]
        headers += [f">{max_reuse_column} %"]
        worksheet.append(headers)

        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        for record in sorted(records, key=lambda item: natural_sort_key(item["trace"])):
            counts = record["reuse_counts"].get(cache_name)
            if counts is not None:
                worksheet.append(build_reuse_row(record["trace"], counts, max_reuse_column))

        first_percentage_column = max_reuse_column + 4
        for row in worksheet.iter_rows(min_row=2, min_col=first_percentage_column):
            for cell in row:
                cell.number_format = "0.0000"

        worksheet.freeze_panes = "B2"
        worksheet.column_dimensions["A"].width = 45
        for column_index in range(2, len(headers) + 1):
            worksheet.column_dimensions[worksheet.cell(1, column_index).column_letter].width = 12

    output_directory = os.path.dirname(output_file)
    if output_directory:
        os.makedirs(output_directory, exist_ok=True)
    workbook.save(output_file)


def main():
    parser = argparse.ArgumentParser(
        description="Extract DEMAND LINE REUSE COUNT histograms from ChampSim output files."
    )
    parser.add_argument("results_dir", nargs="?", default=DEFAULT_RESULTS_DIR)
    parser.add_argument("-o", "--output", default=DEFAULT_OUTPUT_FILE)
    parser.add_argument("--cache-header", default="inc/cache.h")
    args = parser.parse_args()

    if not os.path.isdir(args.results_dir):
        raise SystemExit(f"Results directory not found: {args.results_dir}")

    records = collect_records(args.results_dir)
    write_workbook(records, args.output, load_cache_ways(args.cache_header))
    print(f"Wrote {len(records)} traces to {args.output}")


if __name__ == "__main__":
    main()
