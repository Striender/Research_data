#!/usr/bin/env python3
"""
Extract MSHR Full Load counts and calculate MSHR Full Load per Kilo-Instruction (/ KI)
from ChampSim output logs across cache levels, with consolidated summary comparison sheets.
"""

import argparse
import os
import re
import sys
import warnings
from collections import defaultdict

# Suppress openpyxl warning for long sheet titles
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Auto-detect openpyxl in local virtual environment if not present in system python
try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    script_dir = os.path.dirname(os.path.abspath(__file__))
    candidate_paths = [
        os.path.normpath(os.path.join(script_dir, "..", "myenv", "lib", "python3.12", "site-packages")),
        os.path.normpath(os.path.join(script_dir, "..", "myenv", "lib", "python3.10", "site-packages")),
        os.path.normpath(os.path.join(script_dir, "..", "myenv", "lib", "python3.11", "site-packages")),
    ]
    for p in candidate_paths:
        if os.path.isdir(p) and p not in sys.path:
            sys.path.insert(0, p)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise SystemExit(
            "The 'openpyxl' library is required. Install it with: pip install openpyxl"
        )

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

DEFAULT_RESULTS_DIR = os.path.normpath(
    os.path.join(SCRIPT_DIR, "..", "results_bingo", "cache_seni_test1")
)

DEFAULT_OUTPUT_FILE = os.path.normpath(
    os.path.join(SCRIPT_DIR, "..", "Excel_Output", "Cache_sensi_test2", "mshr_full_load_stats.xlsx")
)

DEFAULT_CACHES = ("L1D", "L2C", "LLC")

# Styling Definitions
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")  # Navy
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
SECTION_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid") # Soft Blue
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="1F4E79")
TITLE_FONT = Font(name="Calibri", size=13, bold=True, color="000000")
BOLD_FONT = Font(name="Calibri", size=11, bold=True)
REGULAR_FONT = Font(name="Calibri", size=11)

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")


def natural_sort_key(value):
    """Sort strings containing numbers in human order (e.g., 16KB, 32KB, 128KB)."""
    return [
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", str(value))
    ]


def parse_file(filepath):
    """
    Efficiently extract MSHR Full Load and instruction counts line-by-line.
    ChampSim prints counter values in standard decimal (std::dec).
    """
    mshr_load_counts = {}
    instructions = None

    try:
        with open(filepath, "r", errors="ignore") as f:
            for line in f:
                # Capture committed instruction count from heartbeat / finished lines
                if "instructions:" in line:
                    match_inst = re.search(r"\binstructions:\s+(\d+)\b", line)
                    if match_inst:
                        instructions = int(match_inst.group(1))

                # Capture MSHR FULL LOAD counter
                if "MSHR FULL" in line and "LOAD:" in line:
                    match_mshr = re.search(r"(\S+)\s+MSHR FULL.*?LOAD:\s+(\d+)", line)
                    if match_mshr:
                        cache_name = match_mshr.group(1).upper()
                        mshr_load_counts[cache_name] = int(match_mshr.group(2))
    except Exception as e:
        print(f"Warning: Failed to read {filepath}: {e}", file=sys.stderr)

    return mshr_load_counts, instructions


def collect_records(results_dir, requested_caches):
    """
    Recursively scans results_dir and organizes records by folder configuration.
    """
    folder_records = {}
    requested_set = {c.upper() for c in requested_caches}

    for root, directories, files in os.walk(results_dir):
        directories.sort(key=natural_sort_key)
        records = []

        for filename in sorted(files, key=natural_sort_key):
            filepath = os.path.join(root, filename)
            if not os.path.isfile(filepath):
                continue

            mshr_load_counts, instructions = parse_file(filepath)

            # Filter only requested caches
            filtered_counts = {
                c: count for c, count in mshr_load_counts.items()
                if c in requested_set
            }

            if not filtered_counts:
                continue

            records.append({
                "trace": filename,
                "instructions": instructions,
                "mshr_full_load": filtered_counts,
            })

        if records:
            relative_folder = os.path.relpath(root, results_dir)
            folder_name = (
                os.path.basename(os.path.abspath(results_dir))
                if relative_folder == "."
                else relative_folder
            )
            folder_records[folder_name] = records

    return folder_records


def make_unique_sheet_name(folder_name, used_names):
    """Create an Excel sheet name with the complete folder/configuration name."""
    name = folder_name.replace(os.sep, "_")
    name = re.sub(r"[:\\/*?\[\]]", "_", name).strip()
    if not name:
        name = "Results"

    base_name = name
    counter = 1
    while name in used_names:
        name = f"{base_name}_{counter}"
        counter += 1

    used_names.add(name)
    return name


def write_summary_sheet(worksheet, folder_records, requested_caches, metric_mode="ki"):
    """
    Builds a consolidated summary sheet comparing either:
    - 'ki': MSHR Full Load / KI
    - 'count': Raw MSHR Full Load count
    across all folders/configurations side-by-side for each cache level.
    """
    sorted_folders = sorted(folder_records.keys(), key=natural_sort_key)

    is_ki = (metric_mode == "ki")
    sheet_title = (
        "Summary: MSHR Full Load per Kilo-Instruction (/ KI)"
        if is_ki else "Summary: MSHR Full Load Count"
    )

    current_row = 1
    worksheet.append([sheet_title])
    title_cell = worksheet.cell(row=current_row, column=1)
    title_cell.font = TITLE_FONT
    current_row += 2

    # Collect all unique traces across all folders
    all_traces = set()
    for records in folder_records.values():
        for r in records:
            all_traces.add(r["trace"])
    sorted_traces = sorted(all_traces, key=natural_sort_key)

    for cache_index, cache_name in enumerate(requested_caches):
        if cache_index > 0:
            current_row += 2  # Spacing between tables

        section_title = f"{cache_name} - MSHR Full Load / KI" if is_ki else f"{cache_name} - MSHR Full Load Count"
        worksheet.append([section_title])
        sec_cell = worksheet.cell(row=current_row, column=1)
        sec_cell.font = SECTION_FONT
        current_row += 1

        headers = ["Trace"] + sorted_folders
        worksheet.append(headers)
        header_row_idx = current_row

        for col_idx, h in enumerate(headers, 1):
            cell = worksheet.cell(row=header_row_idx, column=col_idx)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
        current_row += 1

        # Build data rows
        folder_trace_map = defaultdict(dict)
        for folder, records in folder_records.items():
            for r in records:
                instructions = r["instructions"]
                full_load = r["mshr_full_load"].get(cache_name)
                if full_load is not None:
                    if is_ki:
                        val = (full_load * 1000.0 / instructions) if instructions else None
                    else:
                        val = full_load
                    folder_trace_map[folder][r["trace"]] = val

        start_data_row = current_row
        for trace in sorted_traces:
            row_data = [trace]
            for folder in sorted_folders:
                val = folder_trace_map[folder].get(trace, "")
                row_data.append(val)

            worksheet.append(row_data)
            row_idx = current_row

            # Apply formatting
            for col_idx, val in enumerate(row_data, 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.font = REGULAR_FONT
                cell.border = THIN_BORDER
                if col_idx == 1:
                    cell.alignment = ALIGN_LEFT
                else:
                    cell.alignment = ALIGN_RIGHT
                    if isinstance(val, (int, float)):
                        cell.number_format = "0.0000" if is_ki else "#,##0"
            current_row += 1

        end_data_row = current_row - 1

        # Summary Row: Average across traces
        avg_row = ["Average"]
        for col_idx in range(2, len(headers) + 1):
            col_letter = get_column_letter(col_idx)
            avg_formula = f"=AVERAGE({col_letter}{start_data_row}:{col_letter}{end_data_row})"
            avg_row.append(avg_formula)

        worksheet.append(avg_row)
        for col_idx in range(1, len(headers) + 1):
            cell = worksheet.cell(row=current_row, column=col_idx)
            cell.font = BOLD_FONT
            cell.border = THIN_BORDER
            cell.fill = SECTION_FILL
            if col_idx == 1:
                cell.alignment = ALIGN_LEFT
            else:
                cell.alignment = ALIGN_RIGHT
                cell.number_format = "0.0000" if is_ki else "#,##0"
        current_row += 1

    worksheet.freeze_panes = "B4"
    autofit_columns(worksheet)


def write_folder_detail_sheet(worksheet, folder_name, records, requested_caches):
    """
    Writes detailed tables for each cache level within a specific folder/configuration.
    """
    sorted_records = sorted(records, key=lambda r: natural_sort_key(r["trace"]))

    current_row = 1
    worksheet.append([f"Configuration: {folder_name}"])
    title_cell = worksheet.cell(row=current_row, column=1)
    title_cell.font = TITLE_FONT
    current_row += 2

    for cache_index, cache_name in enumerate(requested_caches):
        if cache_index > 0:
            current_row += 2

        section_title = f"{cache_name} MSHR Full Load Statistics"
        worksheet.append([section_title])
        sec_cell = worksheet.cell(row=current_row, column=1)
        sec_cell.font = SECTION_FONT
        current_row += 1

        headers = ["Trace", "Instructions", "MSHR Full Load", "MSHR Full Load / KI"]
        worksheet.append(headers)
        header_row_idx = current_row

        for col_idx, h in enumerate(headers, 1):
            cell = worksheet.cell(row=header_row_idx, column=col_idx)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
        current_row += 1

        start_data_row = current_row
        for r in sorted_records:
            instructions = r["instructions"]
            full_load = r["mshr_full_load"].get(cache_name)
            rate_ki = (
                (full_load * 1000.0 / instructions)
                if (full_load is not None and instructions)
                else ""
            )

            row_data = [
                r["trace"],
                instructions if instructions is not None else "",
                full_load if full_load is not None else "",
                rate_ki,
            ]
            worksheet.append(row_data)
            row_idx = current_row

            for col_idx, val in enumerate(row_data, 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.font = REGULAR_FONT
                cell.border = THIN_BORDER
                if col_idx == 1:
                    cell.alignment = ALIGN_LEFT
                elif col_idx == 2:
                    cell.alignment = ALIGN_RIGHT
                    if isinstance(val, (int, float)):
                        cell.number_format = "#,##0"
                elif col_idx == 3:
                    cell.alignment = ALIGN_RIGHT
                    if isinstance(val, (int, float)):
                        cell.number_format = "#,##0"
                elif col_idx == 4:
                    cell.alignment = ALIGN_RIGHT
                    if isinstance(val, (int, float)):
                        cell.number_format = "0.0000"
            current_row += 1

        end_data_row = current_row - 1

        # Average Row
        avg_row = [
            "Average",
            f"=AVERAGE(B{start_data_row}:B{end_data_row})",
            f"=AVERAGE(C{start_data_row}:C{end_data_row})",
            f"=AVERAGE(D{start_data_row}:D{end_data_row})",
        ]
        worksheet.append(avg_row)
        for col_idx in range(1, 5):
            cell = worksheet.cell(row=current_row, column=col_idx)
            cell.font = BOLD_FONT
            cell.border = THIN_BORDER
            cell.fill = SECTION_FILL
            if col_idx == 1:
                cell.alignment = ALIGN_LEFT
            elif col_idx in (2, 3):
                cell.alignment = ALIGN_RIGHT
                cell.number_format = "#,##0"
            elif col_idx == 4:
                cell.alignment = ALIGN_RIGHT
                cell.number_format = "0.0000"
        current_row += 1

    worksheet.freeze_panes = "B4"
    autofit_columns(worksheet)


def autofit_columns(worksheet):
    """Adjusts column widths based on maximum contents."""
    for col in worksheet.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or "")
            if val_str.startswith("="):
                max_len = max(max_len, 12)
            else:
                max_len = max(max_len, len(val_str))
        worksheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 52)


def generate_workbook(folder_records, output_file, requested_caches):
    """Builds and saves the complete Excel workbook."""
    wb = Workbook()
    wb.remove(wb.active)  # Remove default blank sheet

    # 1. Primary Summary Sheet: MSHR Full Load / KI comparison
    summary_ki_sheet = wb.create_sheet("Summary_Full_Load_KI")
    write_summary_sheet(summary_ki_sheet, folder_records, requested_caches, metric_mode="ki")

    # 2. Secondary Summary Sheet: MSHR Full Load raw counts
    summary_count_sheet = wb.create_sheet("Summary_Full_Load_Count")
    write_summary_sheet(summary_count_sheet, folder_records, requested_caches, metric_mode="count")

    # 3. Individual Configuration Sheets
    used_sheet_names = {"Summary_Full_Load_KI", "Summary_Full_Load_Count"}
    for folder_name in sorted(folder_records.keys(), key=natural_sort_key):
        sheet_title = make_unique_sheet_name(folder_name, used_sheet_names)
        ws = wb.create_sheet(sheet_title)
        write_folder_detail_sheet(ws, folder_name, folder_records[folder_name], requested_caches)

    os.makedirs(os.path.dirname(os.path.abspath(output_file)), exist_ok=True)
    wb.save(output_file)
    print(f"\n[SUCCESS] Wrote workbook to: {output_file}")


def main():
    parser = argparse.ArgumentParser(
        description="Extract MSHR Full Load and MSHR Full Load / KI with consolidated summary comparison sheets."
    )
    parser.add_argument(
        "results_dir",
        nargs="?",
        default=DEFAULT_RESULTS_DIR,
        help=f"Directory containing simulation output logs (default: {DEFAULT_RESULTS_DIR})",
    )
    parser.add_argument(
        "-o",
        "--output",
        default=DEFAULT_OUTPUT_FILE,
        help=f"Output Excel file path (default: {DEFAULT_OUTPUT_FILE})",
    )
    parser.add_argument(
        "--cache",
        action="append",
        default=None,
        help="Cache levels to extract, e.g. --cache L1D --cache L2C (default: L1D, L2C, LLC)",
    )

    args = parser.parse_args()
    requested_caches = args.cache if args.cache is not None else list(DEFAULT_CACHES)

    if not os.path.isdir(args.results_dir):
        raise SystemExit(f"Error: Results directory not found: {args.results_dir}")

    print(f"Scanning directory: {args.results_dir}")
    print(f"Target caches: {', '.join(requested_caches)}")

    folder_records = collect_records(args.results_dir, requested_caches)

    if not folder_records:
        raise SystemExit(f"No matching MSHR FULL LOAD records found in {args.results_dir}")

    total_files = sum(len(r) for r in folder_records.values())
    print(f"Found {total_files} result files across {len(folder_records)} configuration folders.")

    generate_workbook(folder_records, args.output, requested_caches)


if __name__ == "__main__":
    main()
