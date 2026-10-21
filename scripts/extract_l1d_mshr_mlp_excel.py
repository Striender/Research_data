#!/usr/bin/env python3
"""
Extract L1D MSHR load occupancy cycle breakdown and stats into an Excel file.

Parses ChampSim files for:
  - Non-Zero Occupancy Cycles (Count of cycles where L1D MSHR load occupancy > 0)
  - Non-Zero Occupancy Sum (Sum of MSHR occupancy across all active cycles)
  - L1D Average MSHR Load Occupancy
  - MSHR Extracted Numerator & Denominator
"""

import os
import re
import sys
import argparse
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit(
        "The 'openpyxl' library is required. Install it with: pip install openpyxl"
    )

# --- DEFAULT CONFIGURATION ---
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR

# Supports running directly from PROJECT_ROOT or from scripts/
if not (PROJECT_ROOT / "results_bingo").exists() and (PROJECT_ROOT.parent / "results_bingo").exists():
    PROJECT_ROOT = PROJECT_ROOT.parent

DEFAULT_RESULTS_DIR = PROJECT_ROOT / "results_bingo/cache_seni_test1/checking_MLP/512KB"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "Excel_Output" / "Cache_sensi_test2"
DEFAULT_EXCEL_NAME = "L1D_MSHR_MLP_occupancy_512KB.xlsx"
# ------------------------------


def natural_sort_key(value):
    return [
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", str(value))
    ]


def parse_file(filepath):
    """
    Parses a ChampSim output file to:
      1. Count non-zero occupancy cycles and their sum from the sequence
         preceding 'Heartbeat CPU ...' lines.
      2. Extract 'L1D Average MSHR Load Occupancy' value and the numerator / denominator.
    """
    non_zero_count = 0
    non_zero_sum = 0
    has_stream = False
    l1d_occupancy = None
    mshr_num = None
    mshr_den = None

    with open(filepath, "r", errors="ignore") as f:
        for line in f:
            if "Heartbeat CPU" in line:
                idx = line.find("Heartbeat CPU")
                val_stream = line[:idx]
                tokens = val_stream.split()
                if tokens:
                    has_stream = True
                    for tok in tokens:
                        try:
                            v = int(tok)
                            if v != 0:
                                non_zero_count += 1
                                non_zero_sum += v
                        except ValueError:
                            pass
            elif "L1D Average MSHR Load Occupancy:" in line:
                try:
                    parts = line.strip().split(":")
                    l1d_occupancy = float(parts[1].strip())
                    frac_line = next(f).strip()
                    n_str, d_str = frac_line.split("/")
                    mshr_num = int(n_str.strip())
                    mshr_den = int(d_str.strip())
                except Exception:
                    pass

    # If file did not have the per-cycle stream printed before heartbeat,
    # fallback to the exact MSHR denominator & numerator reported by ChampSim
    if not has_stream and mshr_den is not None and mshr_num is not None:
        non_zero_count = mshr_den
        non_zero_sum = mshr_num

    return {
        "non_zero_count": non_zero_count,
        "non_zero_sum": non_zero_sum,
        "l1d_avg_mshr_occupancy": l1d_occupancy,
        "mshr_numerator": mshr_num,
        "mshr_denominator": mshr_den,
        "has_stream": has_stream,
    }


def make_unique_sheet_name(folder_name, used_names):
    clean_name = re.sub(r"[:\\/*?\[\]]", "_", folder_name.replace("/", "_"))
    sheet_name = clean_name[:31] or "Results"
    original_name = sheet_name
    suffix_idx = 1
    while sheet_name.lower() in used_names:
        suffix = f"_{suffix_idx}"
        sheet_name = original_name[: 31 - len(suffix)] + suffix
        suffix_idx += 1
    used_names.add(sheet_name.lower())
    return sheet_name


def write_excel_workbook(folder_records, output_path):
    wb = Workbook()
    wb.remove(wb.active)  # Remove default active sheet

    header_font = Font(bold=True, color="000000")
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")  # Soft blue
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )

    headers = [
        "Trace",
        "Non-Zero Cycles (Count)",
        "Sum of Non-Zero Occupancy",
        "L1D Average MSHR Load Occupancy",
        "MSHR Numerator",
        "MSHR Denominator",
        "Match Verification",
    ]

    used_sheet_names = set()

    for folder_name in sorted(folder_records.keys(), key=natural_sort_key):
        records = folder_records[folder_name]
        if not records:
            continue

        sheet_title = make_unique_sheet_name(folder_name, used_sheet_names)
        ws = wb.create_sheet(title=sheet_title)

        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        for rec in sorted(records, key=lambda x: natural_sort_key(x["trace"])):
            trace_name = rec["trace"]
            m = rec["metrics"]
            match_str = (
                "MATCH"
                if (m["non_zero_count"] == m["mshr_denominator"] and m["non_zero_sum"] == m["mshr_numerator"])
                else "MISMATCH"
            )

            row_data = [
                trace_name,
                m["non_zero_count"],
                m["non_zero_sum"],
                m["l1d_avg_mshr_occupancy"],
                m["mshr_numerator"],
                m["mshr_denominator"],
                match_str,
            ]
            ws.append(row_data)

            row_idx = ws.max_row
            ws.cell(row=row_idx, column=1).alignment = left_align
            for c in range(2, 7):
                ws.cell(row=row_idx, column=c).alignment = right_align
            ws.cell(row=row_idx, column=7).alignment = center_align

            for c in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=c).border = thin_border

        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(
        description="Extract L1D MSHR cycle count, sum, and occupancy into Excel."
    )
    parser.add_argument(
        "--results-dir",
        type=Path,
        default=DEFAULT_RESULTS_DIR,
        help=f"Path to results directory (default: {DEFAULT_RESULTS_DIR})",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help=f"Path to output directory (default: {DEFAULT_OUTPUT_DIR})",
    )
    parser.add_argument(
        "--filename",
        type=str,
        default=DEFAULT_EXCEL_NAME,
        help=f"Excel file name (default: {DEFAULT_EXCEL_NAME})",
    )
    args = parser.parse_args()

    results_dir = args.results_dir.resolve()
    output_dir = args.output_dir.resolve()
    output_file = output_dir / args.filename

    if not results_dir.is_dir():
        print(f"Error: Results directory '{results_dir}' not found.")
        sys.exit(1)

    print(f"Scanning directory: {results_dir}")
    print(f"Target Excel output: {output_file}")

    folder_records = {}
    total_files = 0

    for root, _, files in os.walk(results_dir):
        if not files:
            continue

        valid_files = [f for f in files if not f.startswith(".")]
        if not valid_files:
            continue

        records = []
        for filename in sorted(valid_files, key=natural_sort_key):
            filepath = os.path.join(root, filename)
            metrics = parse_file(filepath)
            if metrics["l1d_avg_mshr_occupancy"] is not None or metrics["non_zero_count"] > 0:
                records.append({
                    "trace": filename,
                    "metrics": metrics,
                })
                total_files += 1

        if records:
            rel_folder = os.path.relpath(root, results_dir)
            folder_name = results_dir.name if rel_folder == "." else rel_folder
            folder_records[folder_name] = records
            print(f" -> Found {len(records)} records in: {folder_name}")

    if not folder_records:
        print("No valid ChampSim result records found.")
        sys.exit(0)

    print(f"\nWriting Excel workbook with {len(folder_records)} sheet(s)...")
    write_excel_workbook(folder_records, output_file)
    print(f"Successfully generated: {output_file}")
    print(f"Total processed files: {total_files}")


if __name__ == "__main__":
    main()
