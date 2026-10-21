#!/usr/bin/env python3
"""
Extract MPKI metrics for all TLB structures (ITLB, DTLB, STLB, PSCL2-5)
from ChampSim simulation output files into a formatted Excel workbook.
"""

import os
import re
import sys
import json
import csv
import shutil
import copy
import argparse
from collections import defaultdict
from numbers import Number

# Check for openpyxl
try:
    from openpyxl import load_workbook, Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.cell import MergedCell
except ImportError:
    print("The 'openpyxl' library is required to write formatted Excel files.")
    print("Please install it on your server by running: pip install openpyxl")
    sys.exit(1)

# ANSI terminal colors
class TColors:
    HEADER = '\033[95m'
    OKBLUE = '\033[94m'
    OKGREEN = '\033[92m'
    WARNING = '\033[93m'
    FAIL = '\033[91m'
    ENDC = '\033[0m'
    BOLD = '\033[1m'

# Default configuration paths (as requested by user)
DEFAULT_RESULTS_DIR = "../results/SMS_region/testing/16k_16w_4r"
DEFAULT_OUTPUT_DIR = "../Excel_Output/"
EXCEL_OUTPUT_FILE = "TLB_MPKI.xlsx"

PARSER_CACHE_VERSION = 1
FULL_NUMBER_PATTERN = r"([+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?)"
FULL_PRECISION_FORMAT = '0.0000'

def natural_sort_key(s):
    """Sort key that handles numbers naturally (e.g., exp1, exp2, exp10)."""
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'([0-9]+)', str(s))]

def load_json_data(file_path):
    """Loads a JSON file if it exists."""
    if not os.path.exists(file_path):
        return {}
    try:
        with open(file_path, 'r') as f:
            return json.load(f)
    except (IOError, json.JSONDecodeError):
        return {}

def save_json_data(file_path, data):
    """Saves data to a JSON file."""
    try:
        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
        with open(file_path, 'w') as f:
            json.dump(data, f, indent=4)
    except IOError as e:
        print(f"Warning: Could not save cache file at {file_path}: {e}")

def parse_tlb_metrics(filepath):
    """
    Parses a ChampSim output file to extract all TLB MPKI and translation metrics.
    """
    ordered_columns = [
        "Trace File",
        "Experiment",
        "IPC",
        "Instructions",
        "ITLB Total MPKI",
        "DTLB Total MPKI",
        "DTLB Load MPKI",
        "STLB Total MPKI",
        "STLB Load MPKI",
        "PSCL5 Total MPKI",
        "PSCL4 Total MPKI",
        "PSCL3 Total MPKI",
        "PSCL2 Total MPKI",
        "TLB Stall Cycles",
        "TLB Stall %",
        "TLB CPI Contrib"
    ]
    
    if not filepath:
        return ordered_columns

    metrics = {col: None for col in ordered_columns}
    metrics["Trace File"] = os.path.basename(filepath)

    try:
        with open(filepath, 'r', errors='ignore') as f:
            content = f.read()

        # Instructions & IPC
        m_inst = re.search(r"Simulated Instructions\s*:\s*(\d+)", content)
        if m_inst:
            metrics["Instructions"] = int(m_inst.group(1))

        m_ipc = re.search(r"Base CPI / IPC\s*:\s*[\d\.]+\s*/\s*" + FULL_NUMBER_PATTERN, content)
        if m_ipc:
            metrics["IPC"] = float(m_ipc.group(1))
        else:
            # Fallback for heartbeat or standard champsim summary
            m_ipc2 = re.search(r"CPU 0 cumulative IPC:\s*" + FULL_NUMBER_PATTERN, content)
            if m_ipc2:
                metrics["IPC"] = float(m_ipc2.group(1))

        # ITLB Total MPKI
        m_itlb = re.search(r"ITLB TOTAL\s+ACCESS:\s+\d+\s+HIT:\s+\d+\s+MISS:\s+\d+.*?MPKI:\s*" + FULL_NUMBER_PATTERN, content)
        if m_itlb:
            metrics["ITLB Total MPKI"] = float(m_itlb.group(1))

        # DTLB Total MPKI
        m_dtlb = re.search(r"DTLB TOTAL\s+ACCESS:\s+\d+\s+HIT:\s+\d+\s+MISS:\s+\d+.*?MPKI:\s*" + FULL_NUMBER_PATTERN, content)
        if m_dtlb:
            metrics["DTLB Total MPKI"] = float(m_dtlb.group(1))

        # DTLB Load MPKI
        m_dtlb_load = re.search(r"DTLB LOAD TRANSLATION ACCESS:\s+\d+\s+HIT:\s+\d+\s+MISS:\s+\d+.*?MPKI:\s*" + FULL_NUMBER_PATTERN, content)
        if m_dtlb_load:
            metrics["DTLB Load MPKI"] = float(m_dtlb_load.group(1))

        # STLB Total MPKI
        m_stlb = re.search(r"STLB TOTAL\s+ACCESS:\s+\d+\s+HIT:\s+\d+\s+MISS:\s+\d+.*?MPKI:\s*" + FULL_NUMBER_PATTERN, content)
        if m_stlb:
            metrics["STLB Total MPKI"] = float(m_stlb.group(1))

        # STLB Load MPKI
        m_stlb_load = re.search(r"STLB LOAD TRANSLATION ACCESS:\s+\d+\s+HIT:\s+\d+\s+MISS:\s+\d+.*?MPKI:\s*" + FULL_NUMBER_PATTERN, content)
        if m_stlb_load:
            metrics["STLB Load MPKI"] = float(m_stlb_load.group(1))

        # PSCL 5, 4, 3, 2 MPKI
        for lvl in [5, 4, 3, 2]:
            m_pscl = re.search(rf"PSCL{lvl}\s+TOTAL\s+ACCESS:\s+\d+\s+HIT:\s+\d+\s+MISS:\s+\d+.*?MPKI:\s*" + FULL_NUMBER_PATTERN, content)
            if m_pscl:
                metrics[f"PSCL{lvl} Total MPKI"] = float(m_pscl.group(1))

        # ROB Head Translation Stalls (if present from cycle breakdown)
        m_stall = re.search(r"Address Translation \(TLB/PTW\)\s*:\s*(\d+)\s+([\d\.]+)%\s+([\d\.]+)", content)
        if m_stall:
            metrics["TLB Stall Cycles"] = int(m_stall.group(1))
            metrics["TLB Stall %"] = float(m_stall.group(2))
            metrics["TLB CPI Contrib"] = float(m_stall.group(3))

    except Exception as e:
        print(f"Error parsing file {filepath}: {e}")
        return None

    return metrics

def apply_border_to_range(worksheet, row_range, col_range, border_style):
    """Applies a border to a specified grid range."""
    for row in worksheet.iter_rows(min_row=row_range[0], max_row=row_range[1], min_col=col_range[0], max_col=col_range[1]):
        for cell in row:
            cell.border = border_style

def main():
    parser = argparse.ArgumentParser(description="Extract TLB MPKI from ChampSim outputs to Excel")
    parser.add_argument("--results-dir", default=DEFAULT_RESULTS_DIR, help="Path to input results directory")
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR, help="Path to Excel output directory")
    parser.add_argument("--output-file", default=EXCEL_OUTPUT_FILE, help="Excel output file name")
    parser.add_argument("--force", action="store_true", help="Force re-parsing of all files without using cache")
    args, unknown = parser.parse_known_args()

    # Fallback to positional arguments if user provided them like: script.py [dir1] [dir2]
    if len(unknown) >= 1 and not args.results_dir != DEFAULT_RESULTS_DIR:
        args.results_dir = unknown[0]
    if len(unknown) >= 2 and not args.output_dir != DEFAULT_OUTPUT_DIR:
        args.output_dir = unknown[1]

    results_dir = os.path.abspath(args.results_dir)
    output_dir = os.path.abspath(args.output_dir)
    final_output_path = os.path.join(output_dir, args.output_file)
    processed_log_file = os.path.join(output_dir, ".processed_files_tlb.log")
    data_cache_file = os.path.join(output_dir, ".data_cache_tlb.json")

    os.makedirs(output_dir, exist_ok=True)

    print("=" * 80)
    print("CHAMPSIM TLB MPKI EXCEL EXTRACTION TOOL")
    print("=" * 80)
    print(f"Results Directory : {results_dir}")
    print(f"Output Directory  : {output_dir}")
    print(f"Target Excel File : {final_output_path}")

    if not os.path.isdir(results_dir):
        print(f"{TColors.FAIL}Error: Results directory '{results_dir}' not found.{TColors.ENDC}")
        print("Please verify the directory path.")
        return

    processed_files_log = {} if args.force else load_json_data(processed_log_file)
    cached_data = {} if args.force else load_json_data(data_cache_file)
    
    data_by_group = defaultdict(dict)
    new_files_count = 0
    skipped_files_count = 0
    announced_dirs = set()

    print(f"\nScanning directory tree: '{results_dir}'...")

    for root, dirs, files in os.walk(results_dir):
        # Exclude hidden or non-data files
        valid_files = [f for f in files if not f.startswith('.') and not f.endswith(('.py', '.sh', '.json', '.xlsx', '.csv', '.log', '.md'))]
        if not valid_files:
            continue

        relative_path = os.path.relpath(root, results_dir)
        path_parts = relative_path.split(os.sep)

        # Determine group_key and experiment name
        if relative_path == ".":
            group_key = "Default"
            experiment = "All"
        elif len(path_parts) == 1:
            group_key = path_parts[0]
            experiment = path_parts[0]
        else:
            experiment = path_parts[-1]
            group_key = "_".join(path_parts[:-1])

        # Sanitize group_key for Excel sheet name (max 31 characters, no invalid chars)
        sheet_key = re.sub(r'[:\\/?*\[\]]', '_', group_key)[:31]

        for filename in sorted(valid_files):
            filepath = os.path.join(root, filename)
            try:
                file_mod_time = os.path.getmtime(filepath)
            except OSError:
                continue

            log_entry = processed_files_log.get(filepath)
            if (
                not args.force
                and isinstance(log_entry, dict)
                and log_entry.get("mtime") == file_mod_time
                and log_entry.get("parser_cache_version") == PARSER_CACHE_VERSION
            ):
                if sheet_key in cached_data and filepath in cached_data[sheet_key]:
                    data_by_group[sheet_key][filepath] = cached_data[sheet_key][filepath]
                    skipped_files_count += 1
                    continue

            if root not in announced_dirs:
                print(f"{TColors.OKGREEN} -> Reading files in: {relative_path}{TColors.ENDC}")
                announced_dirs.add(root)

            metrics = parse_tlb_metrics(filepath)
            if metrics:
                metrics["Experiment"] = experiment
                data_by_group[sheet_key][filepath] = metrics
                processed_files_log[filepath] = {
                    "mtime": file_mod_time,
                    "parser_cache_version": PARSER_CACHE_VERSION
                }
                new_files_count += 1

    print(f"\nScan complete. Processed {new_files_count} new/modified files ({skipped_files_count} loaded from cache).")

    if not data_by_group:
        print(f"{TColors.WARNING}No ChampSim result files found in '{results_dir}'.{TColors.ENDC}")
        return

    print(f"\nWriting Excel workbook: {final_output_path}...")
    temp_output_path = os.path.join("/tmp", f"{os.getpid()}_{args.output_file}")

    try:
        book = Workbook()
        # Remove default sheet
        if 'Sheet' in book.sheetnames:
            book.remove(book['Sheet'])

        # Preserve any user custom sheets from previous workbook
        existing_custom_sheets = {}
        if os.path.exists(final_output_path):
            try:
                old_book = load_workbook(final_output_path)
                for sname in old_book.sheetnames:
                    if sname not in data_by_group and sname != "Summary":
                        print(f" -> Preserving custom sheet: {sname}")
                        existing_custom_sheets[sname] = old_book[sname]
            except Exception as e:
                print(f"Warning: Could not preserve existing sheets: {e}")

        # Styles
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        main_header_fill = PatternFill(start_color="FFE4B5", end_color="FFE4B5", fill_type="solid") # Moccasin
        sub_header_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")  # Soft Blue
        summary_header_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid") # Soft Green
        
        bold_font = Font(name="Calibri", size=11, bold=True)
        title_font = Font(name="Calibri", size=13, bold=True)
        regular_font = Font(name="Calibri", size=11)
        
        left_align = Alignment(horizontal='left', vertical='center')
        right_align = Alignment(horizontal='right', vertical='center')
        center_align = Alignment(horizontal='center', vertical='center')

        headers = parse_tlb_metrics(None)
        numeric_cols = [c for c in headers if c not in ["Trace File", "Experiment", "Instructions"]]

        summary_rows = []

        # Process each group as a separate sheet
        for sheet_key in sorted(data_by_group.keys(), key=natural_sort_key):
            records = list(data_by_group[sheet_key].values())
            if not records:
                continue

            # Group records by experiment inside this sheet
            exp_groups = defaultdict(list)
            for r in records:
                exp_groups[r.get("Experiment", "All")].append(r)

            worksheet = book.create_sheet(title=sheet_key)
            
            # Row 1: Group Title Header
            worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
            title_cell = worksheet.cell(row=1, column=1, value=f"TLB MPKI Breakdown — {sheet_key}")
            title_cell.font = title_font
            title_cell.fill = main_header_fill
            title_cell.alignment = center_align
            apply_border_to_range(worksheet, (1, 1), (1, len(headers)), thin_border)
            worksheet.row_dimensions[1].height = 28

            current_row = 2

            # Group by experiment inside the sheet if multiple experiments exist
            for experiment in sorted(exp_groups.keys(), key=natural_sort_key):
                exp_records = sorted(exp_groups[experiment], key=lambda x: natural_sort_key(x.get("Trace File", "")))

                # Sub-header with experiment name
                worksheet.cell(row=current_row, column=1, value=f"Experiment: {experiment} (Workloads: {len(exp_records)})").font = bold_font
                current_row += 1

                # Column Headers
                for c_idx, h_text in enumerate(headers, 1):
                    cell = worksheet.cell(row=current_row, column=c_idx, value=h_text)
                    cell.font = bold_font
                    cell.fill = sub_header_fill
                    cell.alignment = center_align
                apply_border_to_range(worksheet, (current_row, current_row), (1, len(headers)), thin_border)
                worksheet.row_dimensions[current_row].height = 22
                current_row += 1

                # Data Rows
                for rec in exp_records:
                    for c_idx, h_text in enumerate(headers, 1):
                        val = rec.get(h_text)
                        cell = worksheet.cell(row=current_row, column=c_idx)
                        if val is None:
                            cell.value = "-"
                            cell.alignment = center_align
                        elif isinstance(val, (int, float)):
                            cell.value = val
                            cell.alignment = right_align
                            cell.number_format = FULL_PRECISION_FORMAT if isinstance(val, float) else '#,##0'
                        else:
                            cell.value = str(val)
                            cell.alignment = left_align
                        cell.font = regular_font
                    apply_border_to_range(worksheet, (current_row, current_row), (1, len(headers)), thin_border)
                    current_row += 1

                # Average Row
                avg_cell = worksheet.cell(row=current_row, column=1, value="Average")
                avg_cell.font = bold_font
                avg_cell.alignment = left_align
                apply_border_to_range(worksheet, (current_row, current_row), (1, len(headers)), thin_border)

                summary_entry = {
                    "Group": sheet_key,
                    "Experiment": experiment,
                    "Workloads": len(exp_records)
                }

                for c_idx, h_text in enumerate(headers, 1):
                    if h_text in numeric_cols:
                        valid_vals = [r[h_text] for r in exp_records if r.get(h_text) is not None and isinstance(r[h_text], (int, float))]
                        if valid_vals:
                            avg_val = sum(valid_vals) / len(valid_vals)
                            cell = worksheet.cell(row=current_row, column=c_idx, value=avg_val)
                            cell.font = bold_font
                            cell.alignment = right_align
                            cell.number_format = FULL_PRECISION_FORMAT
                            summary_entry[h_text] = avg_val

                summary_rows.append(summary_entry)
                current_row += 2  # Blank line before next experiment

            # Auto-fit column widths
            for col_idx in range(1, worksheet.max_column + 1):
                col_letter = get_column_letter(col_idx)
                max_len = max(len(str(cell.value or '')) for cell in worksheet[col_letter] if not isinstance(cell, MergedCell))
                worksheet.column_dimensions[col_letter].width = max(max_len + 3, 14)

            print(f" -> Generated sheet: {sheet_key} ({len(records)} traces)")

        # Create Summary Sheet
        if summary_rows:
            summary_ws = book.create_sheet(title="Summary", index=0)
            summary_ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(summary_rows[0]))
            stitle = summary_ws.cell(row=1, column=1, value="TLB MPKI Overall Summary (Averages Across Configurations)")
            stitle.font = title_font
            stitle.fill = summary_header_fill
            stitle.alignment = center_align
            summary_ws.row_dimensions[1].height = 28

            sum_headers = list(summary_rows[0].keys())
            for c_idx, h_text in enumerate(sum_headers, 1):
                cell = summary_ws.cell(row=2, column=c_idx, value=h_text)
                cell.font = bold_font
                cell.fill = sub_header_fill
                cell.alignment = center_align
            apply_border_to_range(summary_ws, (2, 2), (1, len(sum_headers)), thin_border)
            summary_ws.row_dimensions[2].height = 22

            for r_idx, srow in enumerate(summary_rows, 3):
                for c_idx, h_text in enumerate(sum_headers, 1):
                    val = srow.get(h_text)
                    cell = summary_ws.cell(row=r_idx, column=c_idx)
                    if isinstance(val, (int, float)):
                        cell.value = val
                        cell.alignment = right_align
                        cell.number_format = FULL_PRECISION_FORMAT if isinstance(val, float) else '#,##0'
                    else:
                        cell.value = str(val)
                        cell.alignment = left_align
                    cell.font = regular_font
                apply_border_to_range(summary_ws, (r_idx, r_idx), (1, len(sum_headers)), thin_border)

            for col_idx in range(1, summary_ws.max_column + 1):
                col_letter = get_column_letter(col_idx)
                max_len = max(len(str(cell.value or '')) for cell in summary_ws[col_letter] if not isinstance(cell, MergedCell))
                summary_ws.column_dimensions[col_letter].width = max(max_len + 3, 15)

            # Also export summary CSV
            csv_path = os.path.join(output_dir, "TLB_MPKI_summary.csv")
            with open(csv_path, 'w', newline='') as f_csv:
                writer = csv.DictWriter(f_csv, fieldnames=sum_headers)
                writer.writeheader()
                writer.writerows(summary_rows)
            print(f" -> Exported summary CSV: {csv_path}")

        # Re-attach user preserved custom sheets
        for sname, old_ws in existing_custom_sheets.items():
            new_ws = book.create_sheet(title=sname)
            for row in old_ws.iter_rows():
                for cell in row:
                    new_ws[cell.coordinate].value = cell.value
                    if cell.has_style:
                        new_ws[cell.coordinate].font = copy.copy(cell.font)
                        new_ws[cell.coordinate].border = copy.copy(cell.border)
                        new_ws[cell.coordinate].fill = copy.copy(cell.fill)
                        new_ws[cell.coordinate].number_format = cell.number_format
                        new_ws[cell.coordinate].alignment = copy.copy(cell.alignment)

        # Save workbook
        book.save(temp_output_path)
        shutil.move(temp_output_path, final_output_path)

        save_json_data(data_cache_file, data_by_group)
        save_json_data(processed_log_file, processed_files_log)

        print(f"\n{TColors.OKGREEN}SUCCESS: TLB MPKI Excel file created at:{TColors.ENDC}")
        print(f"{TColors.BOLD}{final_output_path}{TColors.ENDC}\n")

    except Exception as e:
        print(f"{TColors.FAIL}Error writing Excel file: {e}{TColors.ENDC}")
        if os.path.exists(temp_output_path):
            try:
                os.remove(temp_output_path)
            except OSError:
                pass
        raise e

if __name__ == "__main__":
    main()
